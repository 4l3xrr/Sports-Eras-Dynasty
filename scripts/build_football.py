"""
build_football.py
-----------------
Builds Era Draft's football seasons, 1966-2025, plus offensive line units.

Sources
  1999-2025  nflverse (downloads automatically into data/raw/nflverse/)
  1966-1998  "NFL Football Player Stats" on Kaggle (scraped from Pro-Football-Reference in 2017).
             Download it by hand and put its two CSV files (games_*.csv and profiles_*.csv)
             into data/raw/pfr_kaggle/. If they're missing, football starts at 1999.
  O-Lines    u/gyman122's OL Z-Score database (data/reference/gyman122_OL_Z-Scores.xlsx),
             extended to 2022-2025 with the same method.
  Hall of Fame  Kaggle profiles (inducted through 2017) + data/reference/football_hall_of_fame_extra.csv

Output : public/data/football/<year>.json and index.json

Run from the project folder:
    python scripts/build_football.py
"""

import json
import re
import sys
import unicodedata

import numpy as np
import openpyxl
import pandas as pd

from common import PROJECT, RAW, era_for, fetch, player_record, write_seasons
from ratings import HOF_FLOOR, apply_floor, blend, calibrate, report, to_ovr, zscore

# ------------------------------------------------------------------
# 1. SETTINGS
# ------------------------------------------------------------------
NFLVERSE = RAW / "nflverse"
KAGGLE = RAW / "pfr_kaggle"
REFERENCE = PROJECT / "data" / "reference"
RELEASES = "https://github.com/nflverse/nflverse-data/releases/download"
GAMES_URL = "https://raw.githubusercontent.com/nflverse/nfldata/master/data/games.csv"

FIRST_YEAR, NFLVERSE_FIRST, LAST_YEAR = 1966, 1999, 2025
SACKS_FROM, TACKLES_FROM, TARGETS_FROM, MODERN_DEF_FROM = 1982, 1994, 1992, 1999
PRE_SACK_DEFENDER_FLOOR = 90     # Hall of Fame floor for defenders in seasons with no sack stats
MIN_TEAM_GAMES = 3               # games with a team to be listed on its roster
OL_ERA_FLOORS = [(10, 88), (39, 82)]   # top-10 famous line eras: floor 88; ranks 11-39: floor 82

POSITION_GROUPS = {
    "QB": "QB", "RB": "RB", "HB": "RB", "FB": "RB", "WR": "WR", "FL": "WR", "SE": "WR", "TE": "TE",
    "DE": "DL", "DT": "DL", "NT": "DL", "DL": "DL",
    "LB": "LB", "OLB": "LB", "ILB": "LB", "MLB": "LB", "LOLB": "LB", "ROLB": "LB", "LILB": "LB", "RILB": "LB",
    "CB": "CB", "LCB": "CB", "RCB": "CB", "DB": "CB+S",
    "S": "S", "SS": "S", "FS": "S", "SAF": "S",
    "K": "K", "PK": "K",
    # older Pro-Football-Reference labels
    "E": "WR", "SE": "WR", "FL": "WR", "B": "RB", "TB": "RB", "WB": "RB", "HB-FB": "RB",
    "LE": "DL", "RE": "DL", "LDE": "DL", "RDE": "DL", "LDT": "DL", "RDT": "DL", "MG": "DL", "NG": "DL", "DG": "DL",
    "LLB": "LB", "RLB": "LB", "SLB": "LB", "WLB": "LB", "LB-DE": "LB",
    "LS": "", "LH": "CB", "RH": "CB", "DH": "CB",
}
DEFENSE = {"DL", "LB", "CB", "S"}


def normalize(name):
    text = unicodedata.normalize("NFKD", str(name)).encode("ascii", "ignore").decode().lower()
    text = re.sub(r"\b(jr|sr|ii|iii|iv)\b", "", text)
    return re.sub(r"[^a-z]", "", text)


# ------------------------------------------------------------------
# 2. TEAM NAMES (the names Pro-Football-Reference uses, so O-Line data matches)
# ------------------------------------------------------------------
FIXED = {"ATL": "Atlanta Falcons", "BUF": "Buffalo Bills", "CAR": "Carolina Panthers",
         "CHI": "Chicago Bears", "CIN": "Cincinnati Bengals", "DAL": "Dallas Cowboys",
         "DEN": "Denver Broncos", "DET": "Detroit Lions", "GB": "Green Bay Packers",
         "JAX": "Jacksonville Jaguars", "KC": "Kansas City Chiefs", "MIA": "Miami Dolphins",
         "MIN": "Minnesota Vikings", "NO": "New Orleans Saints", "NYG": "New York Giants",
         "NYJ": "New York Jets", "PHI": "Philadelphia Eagles", "PIT": "Pittsburgh Steelers",
         "SEA": "Seattle Seahawks", "SF": "San Francisco 49ers", "TB": "Tampa Bay Buccaneers",
         "HOUT": "Houston Texans", "BALR": "Baltimore Ravens", "CLE": "Cleveland Browns"}


def team_name(franchise, year):
    """franchise is our stable code; returns the name used that season."""
    if franchise == "NE":
        return "Boston Patriots" if year <= 1970 else "New England Patriots"
    if franchise == "IND":
        return "Baltimore Colts" if year <= 1983 else "Indianapolis Colts"
    if franchise == "ARI":
        return "St. Louis Cardinals" if year <= 1987 else "Phoenix Cardinals" if year <= 1993 else "Arizona Cardinals"
    if franchise == "LA":
        return "St. Louis Rams" if 1995 <= year <= 2015 else "Los Angeles Rams"
    if franchise == "LV":
        return "Los Angeles Raiders" if 1982 <= year <= 1994 else "Las Vegas Raiders" if year >= 2020 else "Oakland Raiders"
    if franchise == "LAC":
        return "San Diego Chargers" if year <= 2016 else "Los Angeles Chargers"
    if franchise == "TEN":
        return "Houston Oilers" if year <= 1996 else "Tennessee Oilers" if year <= 1998 else "Tennessee Titans"
    if franchise == "WAS":
        return "Washington Redskins" if year <= 2019 else "Washington Football Team" if year <= 2021 else "Washington Commanders"
    return FIXED[franchise]


# nflverse code -> our franchise code (nflverse mostly uses today's codes already)
NFLVERSE_CODES = {"STL": "LA", "SD": "LAC", "OAK": "LV", "LAR": "LA", "HOU": "HOUT", "BAL": "BALR"}

# Pro-Football-Reference game-log codes (1966-1998) -> franchise code. Year matters for a few.
PFR_CODES = {"ATL": "ATL", "BUF": "BUF", "CAR": "CAR", "CHI": "CHI", "CIN": "CIN", "CLE": "CLE",
             "DAL": "DAL", "DEN": "DEN", "DET": "DET", "GNB": "GB", "GB": "GB", "JAX": "JAX",
             "KAN": "KC", "KC": "KC", "MIA": "MIA", "MIN": "MIN", "NOR": "NO", "NO": "NO",
             "NYG": "NYG", "NYJ": "NYJ", "PHI": "PHI", "PIT": "PIT", "SEA": "SEA", "SFO": "SF",
             "SF": "SF", "TAM": "TB", "TB": "TB", "BOS": "NE", "NWE": "NE", "NE": "NE",
             "IND": "IND", "CLT": "IND", "PHO": "ARI", "ARI": "ARI", "CRD": "ARI",
             "LAR": "LA", "RAM": "LA", "LA": "LA", "RAI": "LV", "LRD": "LV", "LAI": "LV", "OAK": "LV",
             "SDG": "LAC", "SD": "LAC", "TEN": "TEN", "OTI": "TEN", "WAS": "WAS", "RAV": "BALR"}


def pfr_franchise(code, year):
    code = str(code).upper().strip()
    if code == "STL":
        return "ARI" if year <= 1987 else "LA"
    if code == "BAL":
        return "IND" if year <= 1983 else "BALR"
    if code == "HOU":
        return "TEN" if year <= 1996 else "HOUT"
    return PFR_CODES.get(code)


# ------------------------------------------------------------------
# 3. LOAD 1999-2025 (nflverse, week by week so traded players land on each team)
# ------------------------------------------------------------------
print("Getting nflverse files (first run downloads about 250 MB)...")
games = pd.read_csv(fetch(GAMES_URL, NFLVERSE / "games.csv"))

NFLV = {  # nflverse column -> our column
    "attempts": "pass_att", "completions": "pass_cmp", "passing_yards": "pass_yds",
    "passing_tds": "pass_td", "passing_interceptions": "pass_int", "sacks_suffered": "sacks_taken",
    "sack_yards_lost": "sack_yds", "carries": "rush_att", "rushing_yards": "rush_yds",
    "rushing_tds": "rush_td", "rushing_fumbles_lost": "rush_fl", "targets": "targets",
    "receptions": "rec", "receiving_yards": "rec_yds", "receiving_tds": "rec_td",
    "def_sacks": "def_sacks", "def_tackles_solo": "def_solo", "def_tackle_assists": "def_ast",
    "def_interceptions": "def_int", "def_pass_defended": "def_pd", "def_tackles_for_loss": "def_tfl",
    "def_fumbles_forced": "def_ff", "fg_att": "fg_att", "fg_made": "fg_made",
    "fg_made_50_59": "fg50a", "fg_made_60_": "fg50b", "fg_long": "fg_long",
}

weekly, rosters = [], []
for year in range(NFLVERSE_FIRST, LAST_YEAR + 1):
    w = pd.read_csv(fetch(f"{RELEASES}/stats_player/stats_player_week_{year}.csv",
                          NFLVERSE / f"stats_player_week_{year}.csv"), low_memory=False)
    w = w[w.season_type == "REG"].copy()
    w["season"] = year
    weekly.append(w)
    r = pd.read_csv(fetch(f"{RELEASES}/rosters/roster_{year}.csv", NFLVERSE / f"roster_{year}.csv"), low_memory=False)
    r["season"] = year
    rosters.append(r)
weekly = pd.concat(weekly, ignore_index=True)
rosters = pd.concat(rosters, ignore_index=True)

for col in NFLV:
    if col not in weekly:
        weekly[col] = np.nan
weekly = weekly.rename(columns=NFLV)
weekly["team"] = weekly.team.replace(NFLVERSE_CODES)
stat_cols = list(dict.fromkeys(NFLV.values()))
weekly[stat_cols] = weekly[stat_cols].apply(pd.to_numeric, errors="coerce")
weekly["games"] = 1

modern_team = weekly.groupby(["player_id", "season", "team"])[stat_cols + ["games"]].sum(min_count=1).reset_index()
roster_info = (rosters.dropna(subset=["gsis_id"]).sort_values("season")
               .drop_duplicates(["gsis_id", "season"], keep="last")
               .set_index(["gsis_id", "season"]))
names = weekly.dropna(subset=["player_display_name"]).drop_duplicates(["player_id", "season"]) \
    .set_index(["player_id", "season"])[["player_display_name", "position"]]


def modern_meta(pid, season):
    listed = []
    if (pid, season) in roster_info.index:
        r = roster_info.loc[(pid, season)]
        listed += [r.get("position"), r.get("depth_chart_position")]
        birth = str(r.get("birth_date") or "")[:4]
        ref = r.get("pfr_id") if isinstance(r.get("pfr_id"), str) else None
    else:
        birth, ref = "", None
    if (pid, season) in names.index:
        listed.append(names.loc[(pid, season), "position"])
        name = names.loc[(pid, season), "player_display_name"]
    else:
        name = pid
    return name, listed, birth, ref


modern_team["source"] = "nflverse"

# ------------------------------------------------------------------
# 4. LOAD 1966-1998 (Kaggle / Pro-Football-Reference game logs), if present
# ------------------------------------------------------------------
KAGGLE_COLS = {
    "passing_attempts": "pass_att", "passing_completions": "pass_cmp", "passing_yards": "pass_yds",
    "passing_touchdowns": "pass_td", "passing_interceptions": "pass_int", "passing_sacks": "sacks_taken",
    "passing_sacks_yards_lost": "sack_yds", "rushing_attempts": "rush_att", "rushing_yards": "rush_yds",
    "rushing_touchdowns": "rush_td", "receiving_targets": "targets", "receiving_receptions": "rec",
    "receiving_yards": "rec_yds", "receiving_touchdowns": "rec_td", "defense_sacks": "def_sacks",
    "defense_tackles": "def_solo", "defense_tackle_assists": "def_ast", "defense_interceptions": "def_int",
    "field_goal_attempts": "fg_att", "field_goal_makes": "fg_made",
}

# rglob also looks inside subfolders (unzipping often adds an extra folder layer);
# if there's more than one file, use the biggest
def kaggle_files(pattern):
    return sorted(KAGGLE.rglob(pattern), key=lambda p: p.stat().st_size, reverse=True)


def json_records(path):
    """Reads a huge JSON file one record at a time, so a 1 GB file doesn't need 10 GB of memory.
    Works for a list of records ([{...}, {...}]) and for one record per line."""
    decoder, buf = json.JSONDecoder(), ""
    with open(path, encoding="utf-8") as f:
        while True:
            chunk = f.read(4 << 20)
            buf += chunk
            pos = 0
            while True:
                while pos < len(buf) and buf[pos] in " \t\r\n,[":
                    pos += 1
                if pos >= len(buf) or buf[pos] == "]":
                    break
                try:
                    obj, end = decoder.raw_decode(buf, pos)
                except json.JSONDecodeError:
                    if not chunk:
                        raise
                    break                       # record cut off at the chunk edge: read more
                pos = end
                if isinstance(obj, dict):
                    yield obj
            buf = buf[pos:]
            if not chunk:
                return


def read_kaggle(path, keep=None, wanted=None):
    """CSV or JSON into a table. keep(record) filters records while reading;
    wanted limits which fields are kept (saves memory)."""
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, low_memory=False)
        df.columns = [c.strip().lower() for c in df.columns]
        return df
    rows, first = [], None
    for i, rec in enumerate(json_records(path)):
        rec = {k.strip().lower(): v for k, v in rec.items()}
        if first is None:
            first = rec
            print(f"  {path.name} fields: {sorted(rec)}")
        if keep and not keep(rec):
            continue
        rows.append({k: rec.get(k) for k in wanted} if wanted else rec)
        if i and i % 500_000 == 0:
            print(f"    read {i:,} records...")
    return pd.DataFrame(rows)


def regular_season_games(year):
    """Regular-season length: 14 games until 1977, 16 from 1978; 1982 (strike) 9, 1987 (strike) 15."""
    return {1982: 9, 1987: 15}.get(int(year), 14 if year < 1978 else 16)


def record_year(rec):
    try:
        return FIRST_YEAR <= int(rec.get("year")) < NFLVERSE_FIRST
    except (TypeError, ValueError):
        return False


games_files = kaggle_files("games*.csv") + kaggle_files("games*.json")
profile_files = kaggle_files("profiles*.csv") + kaggle_files("profiles*.json")
old_team = pd.DataFrame()
old_records = {}
kaggle_profiles = pd.DataFrame()
if games_files and profile_files:
    print(f"Reading Kaggle files: {games_files[0].name}, {profile_files[0].name} (the big one takes a few minutes)")
    game_fields = ["player_id", "year", "team", "game_id", "date", "player_team_score", "opponent_score", *KAGGLE_COLS]
    g = read_kaggle(games_files[0], keep=record_year, wanted=game_fields)
    # No game ID in this dataset, but a team plays at most once a day, so team + date marks one game
    if ("game_id" not in g or g.game_id.isna().all()) and "date" in g:
        g["game_id"] = g.team.astype(str) + "|" + g.date.astype(str)
    missing = [c for c in ["player_id", "year", "team", "game_id"] if c not in g or g[c].isna().all()]
    if missing:
        sys.exit(f"The Kaggle games file is missing {missing}. Paste the 'fields:' line above to Claude.")
    g["year"] = pd.to_numeric(g.year, errors="coerce")
    g = g[(g.year >= FIRST_YEAR) & (g.year < NFLVERSE_FIRST)].copy()
    g["player_id"] = g.player_id.astype(str)
    g["franchise"] = [pfr_franchise(c, int(y)) for c, y in zip(g.team, g.year)]
    unknown = sorted(set(g.loc[g.franchise.isna(), "team"].astype(str)))
    if unknown:
        print(f"  WARNING: unknown team codes (rows skipped): {unknown}. Add them to PFR_CODES.")
    g = g.dropna(subset=["franchise"])
    # The game logs include playoff games. Playoffs always come after the regular season,
    # so keep each team's first N games by date, where N is that year's schedule length.
    g["when"] = pd.to_datetime(g.date, errors="coerce")
    team_games = g.drop_duplicates(["franchise", "year", "game_id"])[["franchise", "year", "game_id", "when"]]
    team_games = team_games.sort_values("when")
    team_games["n"] = team_games.groupby(["franchise", "year"]).cumcount() + 1
    regular_ids = set(team_games.loc[team_games.n <= team_games.year.map(regular_season_games), "game_id"])
    before = g.game_id.nunique()
    g = g[g.game_id.isin(regular_ids)].copy()
    print(f"  Dropped {before - g.game_id.nunique()} playoff team-games")
    for col in KAGGLE_COLS:
        if col not in g:
            g[col] = np.nan
    g = g.rename(columns=KAGGLE_COLS)
    num = list(KAGGLE_COLS.values())
    g[num] = g[num].apply(pd.to_numeric, errors="coerce")
    g["games"] = 1
    old_team = g.groupby(["player_id", "year", "franchise"])[num + ["games"]].sum(min_count=1).reset_index()
    old_team = old_team.rename(columns={"year": "season", "franchise": "team"})
    old_team["player_id"] = "pfr" + old_team.player_id.astype(str)
    old_team["source"] = "kaggle"
    # team records from each game's score
    per_game = g.drop_duplicates(["franchise", "game_id"])[["franchise", "year", "player_team_score", "opponent_score"]]
    for (fr, yr), rows in per_game.groupby(["franchise", "year"]):
        s, o = pd.to_numeric(rows.player_team_score, errors="coerce"), pd.to_numeric(rows.opponent_score, errors="coerce")
        w, l, t = int((s > o).sum()), int((s < o).sum()), int((s == o).sum())
        old_records[(fr, int(yr))] = f"{w}-{l}" + (f"-{t}" if t else "")
    kaggle_profiles = read_kaggle(profile_files[0])
    kaggle_profiles["player_id"] = "pfr" + kaggle_profiles.player_id.astype(str)
    kaggle_profiles = kaggle_profiles.drop_duplicates("player_id").set_index("player_id")
    print(f"  {len(old_team)} player-team-seasons from {int(old_team.season.min())}-{int(old_team.season.max())}")
    early = old_team[old_team.season <= 1972].groupby("season").agg(players=("player_id", "nunique"), teams=("team", "nunique"))
    print("  Raw Kaggle coverage, early years: " + ", ".join(
        f"{int(y)}: {r.players} players on {r.teams} teams" for y, r in early.iterrows()))
    checks = [("CHI", 1985, "15-1"), ("MIA", 1972, "14-0"), ("SF", 1984, "15-1"), ("PIT", 1978, "14-2")]
    print("  Record check (ours vs real): " + ", ".join(
        f"{yr} {fr} {old_records.get((fr, yr), '?')} vs {real}" for fr, yr, real in checks))
else:
    print("NOTE: Kaggle files not found in data/raw/pfr_kaggle/, so football starts in 1999.")
    found = [str(p.relative_to(KAGGLE)) for p in KAGGLE.rglob("*") if p.is_file()][:10]
    if found:
        print(f"  (That folder has: {found}. Needed: files starting with games_ and profiles_.)")

# ------------------------------------------------------------------
# 5. ONE ROW PER PLAYER-SEASON (all teams added up), with positions
# ------------------------------------------------------------------
per_team = pd.concat([modern_team, old_team], ignore_index=True)
stat_cols = list(dict.fromkeys(NFLV.values()))
for col in stat_cols:
    if col not in per_team:
        per_team[col] = np.nan
season = per_team.groupby(["player_id", "season"])[stat_cols + ["games"]].sum(min_count=1)
season["source"] = per_team.groupby(["player_id", "season"]).source.first()

# Stats that weren't recorded yet are unknown, not zero
yr = season.index.get_level_values("season")
season.loc[yr < SACKS_FROM, "def_sacks"] = np.nan
season.loc[yr < TACKLES_FROM, ["def_solo", "def_ast"]] = np.nan
season.loc[yr < TARGETS_FROM, "targets"] = np.nan
season.loc[yr < MODERN_DEF_FROM, ["def_pd", "def_tfl", "def_ff", "fg50a", "fg50b", "fg_long", "rush_fl"]] = np.nan

hof_extra = pd.read_csv(REFERENCE / "football_hall_of_fame_extra.csv")
hof_extra["key"] = hof_extra.name.map(normalize)
# optional position group ("DL", "CB|S"...) separates same-named players, e.g. the
# Hall of Fame tackle Bob Brown from Bob Brown the Packers defensive tackle
hof_extra["group"] = hof_extra.get("group", pd.Series(dtype=str)).fillna("").astype(str)
kaggle_hof = set()
kaggle_hof_names = {}


def birth_year(value):
    """Finds the year in "1976-03-24", "March 24, 1976", 1976, etc."""
    m = re.search(r"(18|19|20)\d\d", str(value or ""))
    return m.group(0) if m else ""


def pfr_ref(pid):
    """Kaggle IDs are Pro-Football-Reference IDs (sometimes as a link like /players/M/MannPe00.htm).
    Returns the plain ID, which the photo script can look up; None if it doesn't look like one."""
    raw = str(pid)[3:].rsplit("/", 1)[-1].replace(".htm", "")
    return raw if re.fullmatch(r"[A-Za-z.'\-]{2,6}\d{2}", raw) else None
if len(kaggle_profiles):
    hof_col = kaggle_profiles.get("hof_induction_year")
    if hof_col is not None:
        # blanks, "None" and 0 all mean "not in the Hall"
        inducted = pd.to_numeric(hof_col, errors="coerce")
        kaggle_hof = set(kaggle_profiles.index[inducted.notna() & (inducted > 1900)])
        for pid in kaggle_hof:
            p = kaggle_profiles.loc[pid]
            kaggle_hof_names.setdefault(normalize(p["name"]), set()).add(birth_year(p.get("birth_date")))
        print(f"  Kaggle profiles: {len(kaggle_profiles)} players, {len(kaggle_hof)} Hall of Famers")
        if not kaggle_hof:
            sample = hof_col.dropna().astype(str).value_counts().head(8).to_dict()
            print(f"  hof_induction_year values look like: {sample or 'all empty'}")
            for star in ["Walter Payton", "Joe Montana", "Lawrence Taylor"]:
                hit = kaggle_profiles[kaggle_profiles["name"] == star]
                if len(hit):
                    print(f"  {star}: {hit.iloc[0].drop(['height', 'weight'], errors='ignore').to_dict()}")


def is_hof(name, birth, year, pid, source, groups=()):
    if pid in kaggle_hof:
        return True
    key = normalize(name)
    # Modern (nflverse) players: find Kaggle-listed Hall of Famers by name AND birth year
    if source == "nflverse" and birth and birth in kaggle_hof_names.get(key, ()):
        return True
    match = hof_extra[(hof_extra.key == key) & (hof_extra.first_season <= year) & (hof_extra.last_season >= year)]
    return any(not g or set(g.split("|")) & set(groups) for g in match.group)


def groups_from_stats(row):
    """Older players with no usable position label: guess the group from what they did.
    Offensive linemen have no stats, so they stay out (lines are rated as units)."""
    v = lambda col: 0 if pd.isna(row.get(col)) else row.get(col)
    if v("pass_att") >= 50:
        return ["QB"]
    if v("fg_att") >= 5:
        return ["K"]
    if v("rush_att") >= 30 and v("rush_att") >= v("rec"):
        return ["RB"]
    if v("rec") >= 10:
        return ["WR"]
    if v("def_sacks") >= 2:
        return ["DL"]
    if v("def_int") >= 2:
        return ["CB", "S"]
    return []


inferred, unmapped = 0, {}
meta = []
for (pid, yr_), row in season.iterrows():
    if row.source == "nflverse":
        name, listed, birth, ref = modern_meta(pid, yr_)
    else:
        prof = kaggle_profiles.loc[pid] if pid in kaggle_profiles.index else None
        name = prof["name"] if prof is not None else pid
        # positions come as "DE", "DT-DE", "WR/PR", "LB, DE"...: split on any separator
        listed = [p for p in re.split(r"[,/\-;\s]+", str(prof.get("position") if prof is not None else "").upper()) if p]
        birth, ref = (birth_year(prof.get("birth_date")) if prof is not None else ""), pfr_ref(pid)
    groups = []
    for p in listed:
        for grp in str(POSITION_GROUPS.get(str(p).upper(), "")).split("+"):
            if grp and grp not in groups:
                groups.append(grp)
    if not groups and row.source == "kaggle":
        groups = groups_from_stats(row)
        if groups:
            inferred += 1
        else:
            unmapped[" ".join(listed) or "(blank)"] = unmapped.get(" ".join(listed) or "(blank)", 0) + 1
    meta.append((name, groups, ref, is_hof(name, birth, yr_, pid, row.source, groups)))
season["name"], season["groups"], season["ref"], season["hof"] = zip(*meta)
if inferred or unmapped:
    top = sorted(unmapped.items(), key=lambda kv: -kv[1])[:12]
    print(f"  Kaggle positions: guessed {inferred} player-seasons from their stats; "
          f"left out {sum(unmapped.values())} with no usable position (mostly linemen). Most common: {top}")
# Health check: Hall of Famers on our list (skill and defensive players) who never matched anyone
found = set(season.loc[season.hof, "name"].map(normalize))
first_data_year = int(season.index.get_level_values("season").min())
expected = hof_extra[(hof_extra.last_season >= first_data_year) & ~hof_extra.group.isin(["OL", "P"])]
unmatched = sorted(set(expected.name[~expected.key.isin(found)]))
print(f"  Hall of Famers matched: {len(found)} players. "
      + (f"Not found (name spelled differently?): {unmatched}" if unmatched else "All listed ones found."))
season = season[season.groups.map(len) > 0].copy()
season["group"] = season.groups.map(lambda g: g[0])

# ------------------------------------------------------------------
# 6. WHO MAKES A ROSTER, AND RATINGS BY POSITION
# ------------------------------------------------------------------
c = lambda col: season[col]
yr = season.index.get_level_values("season")
tackles = c("def_solo").fillna(0) + c("def_ast").fillna(0)
tackles = tackles.where(yr >= TACKLES_FROM)
season["tackles"] = tackles
grp = season.group
qualifies = (
    ((grp == "QB") & (c("pass_att") >= 100)) |
    ((grp == "RB") & (c("rush_att").fillna(0) + c("rec").fillna(0) >= 50)) |
    (grp.isin(["WR", "TE"]) & (c("rec") >= 15)) |
    ((grp == "K") & (c("fg_att") >= 8)) |
    (grp.isin(DEFENSE) & (c("games") >= 8) & (
        (yr < TACKLES_FROM) | (tackles >= 15) | (c("def_sacks").fillna(0) >= 2) | (c("def_int").fillna(0) >= 1)))
)
season = season[qualifies].copy()
c = lambda col: season[col]
yr = season.index.get_level_values("season")


def shrink(rate, volume, prior, k):
    return (rate * volume + prior * k) / (volume + k)


anya = (c("pass_yds") + 20 * c("pass_td") - 45 * c("pass_int") - c("sack_yds").fillna(0)) / \
       (c("pass_att") + c("sacks_taken").fillna(0))
ypc = c("rush_yds") / c("rush_att").replace(0, np.nan)
fg_pct = c("fg_made") / c("fg_att").replace(0, np.nan)
season["long_fg"] = c("fg50a") + c("fg50b").fillna(0)

COMPONENTS = {  # position -> [(stat series, weight)]
    "QB": [(shrink(anya, c("pass_att"), 4.5, 150), .4), (c("pass_yds"), .3),
           (c("pass_td") - c("pass_int"), .2), (c("rush_yds"), .1)],
    "RB": [(c("rush_yds").fillna(0) + c("rec_yds").fillna(0), .45), (c("rush_td").fillna(0) + c("rec_td").fillna(0), .25),
           (shrink(ypc.fillna(3.8), c("rush_att").fillna(0), 3.8, 80), .2), (c("rec"), .1)],
    "WR": [(c("rec_yds"), .5), (c("rec_td"), .25), (c("rec"), .25)],
    "TE": [(c("rec_yds"), .5), (c("rec_td"), .25), (c("rec"), .25)],
    "DL": [(c("def_sacks"), .5), (c("def_tfl"), .15), (c("tackles"), .2), (c("def_ff"), .15), (c("games"), .1)],
    "LB": [(c("tackles"), .35), (c("def_sacks"), .25), (c("def_int"), .15), (c("def_pd"), .1),
           (c("def_tfl"), .1), (c("games"), .05)],
    "CB": [(c("def_int"), .35), (c("def_pd"), .25), (c("tackles"), .2), (c("def_ff"), .1), (c("games"), .1)],
    "S":  [(c("def_int"), .35), (c("def_pd"), .25), (c("tackles"), .2), (c("def_ff"), .1), (c("games"), .1)],
    "K":  [(c("fg_made"), .4), (shrink(fg_pct, c("fg_att"), .75, 15), .4), (season.long_fg, .2)],
}

season["z"] = np.nan
for group, parts in COMPONENTS.items():
    rows = season.group == group
    if not rows.any():
        continue
    g_years = yr[rows]
    zs = [zscore(series[rows], g_years).where(series[rows].notna()) for series, _ in parts]
    blended = blend(zs, [w for _, w in parts])
    season.loc[rows, "z"] = zscore(blended, g_years).values
season["z"] = season.z.fillna(0)

# Each position gets its own anchors (and its fair share of 99s). Sack totals have a much
# longer tail than passing stats, so one shared scale would hand most 99s to pass rushers.
n_seasons = LAST_YEAR - (int(yr.min()) if len(yr) else NFLVERSE_FIRST) + 1
season["ovr"] = 0
for group, rows in season.groupby("group"):
    anchors = calibrate(rows.z, n_seasons, share=len(rows) / len(season))
    season.loc[rows.index, "ovr"] = to_ovr(rows.z, anchors).values

max_games = season.groupby(level="season").games.transform("max")
regular = season.games >= max_games / 2
season["limited"] = season.group.isin(DEFENSE) & (yr < SACKS_FROM)
season["ovr"], floored = apply_floor(season.ovr, season.hof, regular, HOF_FLOOR)
season["ovr"], floored_def = apply_floor(season.ovr, season.hof & season.limited, regular, PRE_SACK_DEFENDER_FLOOR)
season["floored"] = floored | floored_def
report("Football", season.ovr, n_seasons)


def stat_line(r):
    i = lambda v: 0 if pd.isna(v) else int(v)
    g = r.group
    if g == "QB":
        return {"YDS": i(r.pass_yds), "TD": i(r.pass_td), "INT": i(r.pass_int)}
    if g == "RB":
        return {"RUSH": i(r.rush_yds), "REC": i(r.rec_yds), "TD": i(r.rush_td) + i(r.rec_td)}
    if g in ("WR", "TE"):
        return {"REC": i(r.rec), "YDS": i(r.rec_yds), "TD": i(r.rec_td)}
    if g == "K":
        return {"FG": f"{i(r.fg_made)}/{i(r.fg_att)}", "PCT": f"{100 * r.fg_made / r.fg_att:.0f}%" if r.fg_att else "-"}
    line = {}
    if pd.notna(r.def_sacks):
        line["SACKS"] = float(r.def_sacks)
    if pd.notna(r.tackles):
        line["TKL"] = i(r.tackles)
    line["INT"] = i(r.def_int)
    line["GP"] = i(r.games)
    return dict(list(line.items())[:3])


# ------------------------------------------------------------------
# 7. OFFENSIVE LINE UNITS
# ------------------------------------------------------------------
wb = openpyxl.load_workbook(REFERENCE / "gyman122_OL_Z-Scores.xlsx", read_only=True, data_only=True)
line_rows = [r for r in wb["Sheet30"].iter_rows(min_row=2, values_only=True) if r[0]]
lines = pd.DataFrame([(int(r[0][:4]), re.sub(r"[*+]", "", r[0][5:]).strip(), r[1], r[2], r[3]) for r in line_rows],
                     columns=["season", "team_name", "combined", "protection", "rushing"])
lines["source"] = "gyman122"

# Famous eras: "1989-2000 Dallas Cowboys" ranked 1-39
eras = []
for r in wb["Best OL Rank by Prime Total (Pr"].iter_rows(min_row=2, values_only=True):
    if r[0] and r[1]:
        m = re.match(r"(\d{4})-(\d{4}) (.+)", str(r[1]))
        if m:
            eras.append((int(r[0]), int(m.group(1)), int(m.group(2)), m.group(3).strip()))

# 2022-2025: same method as his data (checked against his 1999-2021 numbers):
# protection = sack rate, rushing = (rush yds + 20*TD - 45*fumbles lost) per carry,
# each as a z-score against the five seasons around it.
team_year = weekly.groupby(["team", "season"]).agg(
    att=("pass_att", "sum"), sk=("sacks_taken", "sum"), car=("rush_att", "sum"), ry=("rush_yds", "sum"),
    rtd=("rush_td", "sum"), fl=("rush_fl", "sum")).reset_index()
team_year["sack_rate"] = team_year.sk / (team_year.att + team_year.sk)
team_year["rush_eff"] = (team_year.ry + 20 * team_year.rtd - 45 * team_year.fl) / team_year.car


def window_z(col):
    out = pd.Series(index=team_year.index, dtype=float)
    for y in team_year.season.unique():
        pool = team_year[(team_year.season >= y - 2) & (team_year.season <= y + 2)][col]
        rows = team_year.season == y
        out[rows] = (team_year.loc[rows, col] - pool.mean()) / pool.std()
    return out


team_year["protection"] = -window_z("sack_rate")
team_year["rushing"] = window_z("rush_eff")
team_year["combined"] = (team_year.protection + team_year.rushing) / 2
team_year["team_name"] = [team_name(t, y) for t, y in zip(team_year.team, team_year.season)]
newest = int(lines.season.max())
computed = team_year[team_year.season > newest][["season", "team_name", "combined", "protection", "rushing"]].copy()
computed["source"] = "computed"
lines = pd.concat([lines, computed], ignore_index=True)
lines = lines[lines.season >= FIRST_YEAR].copy()
lines["z"] = zscore(lines.combined)
lines["ovr"] = to_ovr(lines.z, calibrate(lines.z, lines.season.nunique()))


def era_floor(name, year):
    best = None
    for rank, start, end, team in eras:
        if team == name and start <= year <= end and (best is None or rank < best[0]):
            best = (rank, start, end)
    if not best:
        return 0, None
    floor = next(f for limit, f in OL_ERA_FLOORS if best[0] <= limit)
    return floor, best


line_lookup = {}
for _, r in lines.iterrows():
    floor, era = era_floor(r.team_name, r.season)
    line_lookup[(r.season, r.team_name)] = (max(int(r.ovr), floor), r, era, floor > r.ovr)

# Starting five, from snap counts (2013 on)
starters = {}
for year in range(2013, LAST_YEAR + 1):
    path = NFLVERSE / f"snap_counts_{year}.csv"
    try:
        snaps = pd.read_csv(fetch(f"{RELEASES}/snap_counts/snap_counts_{year}.csv", path))
    except Exception:
        continue
    snaps = snaps[(snaps.game_type == "REG") & snaps.position.isin(["T", "G", "C", "OT", "OG", "OL"])]
    snaps["team"] = snaps.team.replace(NFLVERSE_CODES)
    top = snaps.groupby(["team", "player"]).offense_snaps.sum().reset_index()
    for team, rows in top.groupby("team"):
        starters[(year, team)] = [n.split()[-1] for n in rows.nlargest(5, "offense_snaps").player]
report("O-Line units", lines.ovr, lines.season.nunique())


def line_unit(franchise, year):
    name = team_name(franchise, year)
    if (year, name) not in line_lookup:
        return None
    ovr, r, era, raised = line_lookup[(year, name)]
    unit = {"id": f"OL-{franchise}-{year}", "name": "Offensive Line", "positions": ["OL"], "ovr": int(ovr),
            "unit": 1, "stats": {"PROTECT": f"{r.protection:+.1f}", "RUSH": f"{r.rushing:+.1f}"}}
    if (year, franchise) in starters:
        unit["members"] = starters[(year, franchise)]
    if era:
        unit["era"] = {"rank": era[0], "span": f"{era[1]}-{era[2]}"}
    return unit


# ------------------------------------------------------------------
# 8. TEAM RECORDS AND WRITING
# ------------------------------------------------------------------
reg_games = games[(games.game_type == "REG") & games.result.notna()].copy()
for col in ("home_team", "away_team"):
    reg_games[col] = reg_games[col].replace(NFLVERSE_CODES)


def modern_record(team, year):
    g = reg_games[reg_games.season == year]
    home, away = g[g.home_team == team], g[g.away_team == team]
    w = (home.result > 0).sum() + (away.result < 0).sum()
    l = (home.result < 0).sum() + (away.result > 0).sum()
    t = (home.result == 0).sum() + (away.result == 0).sum()
    return f"{w}-{l}" + (f"-{t}" if t else "")


rated = season
per_team = per_team[per_team.games >= MIN_TEAM_GAMES]

# 1966-1969 box scores kept almost nothing for defenders, so most of them are missing.
# Hall of Fame defenders from those years are added by hand at the old-defender floor (90),
# marked "Limited stats" so it's clear their rating isn't from numbers.
hof_defenders = {}
hof_def_path = REFERENCE / "football_hof_defenders_1966_1969.csv"
if hof_def_path.exists():
    for r in pd.read_csv(hof_def_path).itertuples():
        for y in range(int(r.first_year), int(r.last_year) + 1):
            hof_defenders.setdefault((y, r.team), []).append((r.name, r.group))
added_defenders = 0


def add_hof_defenders(players, franchise, year):
    global added_defenders
    on_team = {normalize(p["name"]) for p in players}
    for name, group in hof_defenders.get((year, franchise), []):
        if normalize(name) in on_team:
            continue
        players.append(player_record(f"hof-{normalize(name)}-{year}", name, [group], PRE_SACK_DEFENDER_FLOOR,
                                     {"STATS": "not kept"}, hof=True, floored=True, limited=True))
        added_defenders += 1
seasons_out = []
for year, members in per_team.groupby("season"):
    year = int(year)
    out_teams = []
    for franchise, rows in members.groupby("team"):
        players = []
        for pid in rows.player_id:
            if (pid, year) not in rated.index:
                continue
            r = rated.loc[(pid, year)]
            players.append(player_record(f"{pid}-{year}", r["name"], r.groups, r.ovr, stat_line(r),
                                         ref=r.ref if isinstance(r.ref, str) else None, hof=r.hof,
                                         floored=r.floored, limited=r.limited))
        unit = line_unit(franchise, year)
        if unit:
            players.append(unit)
        if len(players) < 8:
            continue
        add_hof_defenders(players, franchise, year)
        players.sort(key=lambda p: -p["ovr"])
        record = modern_record(franchise, year) if year >= NFLVERSE_FIRST else old_records.get((franchise, year), "")
        out_teams.append({"id": f"{franchise}-{year}", "code": franchise, "name": team_name(franchise, year),
                          "record": record, "players": players})
    seasons_out.append({"year": year, "era": era_for(year), "teams": out_teams})

write_seasons("football", seasons_out)
if hof_defenders:
    print(f"  Added {added_defenders} Hall of Fame defender-seasons by hand (1966-1969).")

# How many real teams each early season has (teams with too few players are left out)
REAL_TEAMS = {1966: 24, 1967: 25, 1968: 26, 1969: 26}
for s in seasons_out:
    want = REAL_TEAMS.get(s["year"], 26 if s["year"] < 1976 else 28 if s["year"] < 1995 else 30)
    if s["year"] < NFLVERSE_FIRST and len(s["teams"]) < want:
        sizes = sorted(len(t["players"]) for t in s["teams"])
        print(f"  NOTE: {s['year']} has {len(s['teams'])} of {want} teams (players per team: {sizes})")

# A quick health check that shows up if the old data looks thin
for s in seasons_out:
    if s["year"] < NFLVERSE_FIRST:
        defenders = sum(1 for t in s["teams"] for p in t["players"] if set(p["positions"]) & DEFENSE)
        if defenders < 5 * max(len(s["teams"]), 1):
            print(f"  NOTE: {s['year']} has only {defenders} defenders across {len(s['teams'])} teams.")
